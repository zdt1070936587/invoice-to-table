from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from fapiao.cli import main
from fapiao.ed25519 import public_key, sign
from fapiao.license import canonical_payload
from fapiao.pipeline import run
import json
import base64

DATA = Path(__file__).resolve().parents[1] / "samples" / "data"
BUYER = "91330000MA0000002Y"
ROOT = Path(__file__).resolve().parents[1]


def test_sample_directory_free_limit(tmp_path: Path):
    before = {path: path.stat().st_mtime_ns for path in DATA.rglob("*") if path.is_file()}
    output = tmp_path / "进项_样本.xlsx"
    result = run(DATA, BUYER, "in", output, today=__import__("datetime").date(2026, 10, 8))
    after = {path: path.stat().st_mtime_ns for path in DATA.rglob("*") if path.is_file()}
    assert before == after
    assert result.qualified_count == 5
    assert result.exception_count == 12
    numbers = [invoice.number for invoice in result.invoices]
    assert numbers == [
        "24330000000000000001",
        "24330000000000000002",
        "24330000000000000003",
        "24330000000000000004",
        "24330000000000000005",
    ]
    assert result.invoices[0].total == Decimal("113.00")
    assert "en/invoice.xml" in result.invoices[0].source_files
    assert any(name.endswith("invoice.pdf") for name in result.invoices[0].source_files)
    red = result.invoices[3]
    assert red.total == Decimal("-113.00")
    assert red.is_red
    codes = [row.code for row in result.exceptions]
    assert codes.count("DUP_INVOICE") == 6
    assert codes.count("LICENSE_LIMIT") == 1
    assert codes.count("NEEDS_OCR") == 2
    assert codes.count("OFD_NO_XML") == 1
    assert codes.count("RED_INVOICE") == 1
    assert codes.count("UNPACK_REJECTED") == 1
    assert codes == sorted(codes)

    workbook = load_workbook(output)
    summary = workbook["发票汇总"]
    assert "未验签、不查验" in summary["A1"].value
    assert summary["A2"].value == "发票号码"
    assert summary["I3"].value == 100
    assert summary["K3"].value == 113
    assert summary["I3"].number_format == "0.00"
    lines = workbook["商品明细"]
    assert lines["A2"].value == "项目名称" or lines["C2"].value == "项目名称"
    rate_cell = None
    for row in lines.iter_rows(min_row=3, max_col=9):
        if row[2].value == "咨询服务":
            rate_cell = row[8]
    assert rate_cell is not None
    assert rate_cell.number_format == "0%"
    assert abs(rate_cell.value - 0.03) < 1e-9
    errors = workbook["异常"]
    assert [errors.cell(row, 3).value for row in range(3, 15)] == codes
    limited = [row for row in result.exceptions if row.code == "LICENSE_LIMIT"]
    assert limited[0].number == "24330000000000000006"
    assert limited[0].message == "超过免费额度，未写入汇总"


def test_paid_license_writes_sixth_invoice(tmp_path: Path):
    secret = bytes.fromhex("11" * 32)
    pub = public_key(secret)
    payload = canonical_payload("发票成表", "示例会计", "2027-10-08", 2)
    document = {
        "product": "发票成表",
        "licensee": "示例会计",
        "expiry": "2027-10-08",
        "seats": 2,
        "signature": base64.b64encode(sign(payload, secret)).decode("ascii"),
    }
    (tmp_path / "license.json").write_text(json.dumps(document), encoding="utf-8")
    output = tmp_path / "out.xlsx"
    result = run(
        DATA,
        BUYER,
        "in",
        output,
        license_dir=tmp_path,
        license_public_key=pub,
        today=__import__("datetime").date(2026, 10, 8),
    )
    assert result.qualified_count == 6
    assert all(row.code != "LICENSE_LIMIT" for row in result.exceptions)
    assert result.invoices[-1].number == "24330000000000000006"


def test_cli_run_and_empty_tax_id(tmp_path: Path, capsys):
    output = tmp_path / "cli.xlsx"
    code = main(["run", str(DATA), "--tax-id", BUYER, "--direction", "in", "--out", str(output)])
    captured = capsys.readouterr()
    assert code == 0
    assert "合格 5 张" in captured.out
    assert output.is_file()
    empty = main(["run", str(DATA), "--tax-id", "   ", "--direction", "in", "--out", str(tmp_path / "x.xlsx")])
    assert empty == 2


def test_direction_out_filters_buyer_only_pack(tmp_path: Path):
    result = run(DATA / "en" / "invoice.xml", "91330000MA0000001X", "out", tmp_path / "销项.xlsx")
    assert result.qualified_count == 1
    result_wrong = run(DATA / "en" / "invoice.xml", BUYER, "out", tmp_path / "销项2.xlsx")
    assert result_wrong.qualified_count == 0
    assert result_wrong.exceptions[0].code == "TAX_ID_MISMATCH"


def test_ui_has_no_second_parser():
    source = (ROOT / "src" / "fapiao" / "ui.py").read_text(encoding="utf-8")
    assert "xml_parser" not in source
    assert "ofd_parser" not in source
    assert "pdf_parser" not in source
    assert "pipeline" in source
    assert "文件只在这台电脑处理" in source


def test_unsupported_kind(tmp_path: Path):
    path = tmp_path / "rail.xml"
    path.write_text(
        """<?xml version="1.0" encoding="utf-8"?>
        <EInvoice><Header><EIid>24330000000000000088</EIid></Header>
        <Title>铁路电子客票</Title></EInvoice>
        """,
        encoding="utf-8",
    )
    result = run(path, BUYER, "in", tmp_path / "out.xlsx")
    assert result.qualified_count == 0
    assert result.exceptions[0].code == "UNSUPPORTED_KIND"


def test_unknown_file(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("not an invoice", encoding="utf-8")
    result = run(path, BUYER, "in", tmp_path / "out.xlsx")
    assert result.exceptions[0].code == "UNKNOWN_FILE"
