"""写出发票汇总、商品明细、异常三张表。"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fapiao.models import ExceptionRow, Invoice

DISCLAIMER = "未验签、不查验。结果用于整理，不代替税务局查验。"

SUMMARY_HEADERS = (
    "发票号码",
    "开票日期",
    "票种",
    "是否红字",
    "购买方名称",
    "购买方税号",
    "销售方名称",
    "销售方税号",
    "合计金额",
    "合计税额",
    "价税合计",
    "开票人",
    "来源文件",
    "解析级别",
)
LINE_HEADERS = (
    "发票号码",
    "开票日期",
    "项目名称",
    "规格型号",
    "单位",
    "数量",
    "单价",
    "金额",
    "税率",
    "税额",
    "税收分类编码",
    "编码状态",
)
EXCEPTION_HEADERS = ("文件名", "发票号码", "原因码", "说明", "解析级别")
_MONEY_FORMAT = "0.00"
_RATE_FORMAT = "0%"


def export_workbook(path: Path, invoices: list[Invoice], exceptions: list[ExceptionRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "发票汇总"
    lines = workbook.create_sheet("商品明细")
    errors = workbook.create_sheet("异常")
    _write_summary(summary, invoices)
    _write_lines(lines, invoices)
    _write_exceptions(errors, exceptions)
    workbook.save(path)


def _write_summary(sheet: Worksheet, invoices: list[Invoice]) -> None:
    _prepare(sheet, SUMMARY_HEADERS)
    for row_index, invoice in enumerate(invoices, start=3):
        values = [
            invoice.number,
            invoice.issue_date,
            invoice.kind,
            "是" if invoice.is_red else "否",
            invoice.buyer_name,
            invoice.buyer_tax_id,
            invoice.seller_name,
            invoice.seller_tax_id,
            _number(invoice.amount),
            _number(invoice.tax_amount),
            _number(invoice.total),
            invoice.issuer,
            "; ".join(invoice.source_files),
            invoice.parse_level,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
        for column in (9, 10, 11):
            sheet.cell(row_index, column).number_format = _MONEY_FORMAT
        for column in (1, 6, 8):
            sheet.cell(row_index, column).number_format = "@"


def _write_lines(sheet: Worksheet, invoices: list[Invoice]) -> None:
    _prepare(sheet, LINE_HEADERS)
    row_index = 3
    for invoice in invoices:
        for line in invoice.lines:
            rate_value: str | float
            rate_is_number = line.tax_rate_text not in ("免税", "不征税")
            if rate_is_number:
                rate_value = _number(line.tax_rate)
            else:
                rate_value = line.tax_rate_text
            values = [
                invoice.number,
                invoice.issue_date,
                line.name,
                line.spec,
                line.unit,
                _optional_number(line.quantity),
                _optional_number(line.unit_price),
                _number(line.amount),
                rate_value,
                _number(line.tax_amount),
                line.tax_code,
                line.tax_code_status,
            ]
            for column, value in enumerate(values, start=1):
                if value is not None:
                    sheet.cell(row_index, column, value)
            sheet.cell(row_index, 1).number_format = "@"
            sheet.cell(row_index, 11).number_format = "@"
            for column in (6, 7, 8, 10):
                if sheet.cell(row_index, column).value is not None:
                    sheet.cell(row_index, column).number_format = _MONEY_FORMAT
            if rate_is_number:
                sheet.cell(row_index, 9).number_format = _RATE_FORMAT
            row_index += 1


def _write_exceptions(sheet: Worksheet, exceptions: list[ExceptionRow]) -> None:
    _prepare(sheet, EXCEPTION_HEADERS)
    for row_index, item in enumerate(exceptions, start=3):
        values = [item.filename, item.number, item.code, item.message, item.parse_level]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
        sheet.cell(row_index, 2).number_format = "@"


def _prepare(sheet: Worksheet, headers: tuple[str, ...]) -> None:
    sheet["A1"] = DISCLAIMER
    sheet["A1"].font = Font(bold=True)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        cell.font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(column)].width = max(14, len(header) + 4)
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"


def _number(value: Decimal) -> float:
    return float(value)


def _optional_number(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value)
